#!/usr/bin/env python3
"""Hourly sync for the almentor PMO Recognition dashboard.

1. Downloads PMO_Recognition_Sheet_2026.xlsx from the SharePoint share link
   (secret SHEET_URL), reads the "For Business Review" tab and parses it.
2. Applies project-manager changes sent from the dashboard (env CHANGES).
3. Writes encrypted data for the site (AES-256-GCM with secret DATA_KEY):
     data/snapshot.enc, data/state.enc   -> committed to main (history = audit trail)
     site/data.enc, site/status.json     -> published to GitHub Pages
If the download or validation fails, the last good snapshot is kept and the
failure is reported in status.json so the dashboard can show it.
"""
import base64, datetime as dt, hashlib, io, json, os, re, shutil, sys, traceback

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheetparse import parse_grid, check, SHEET  # noqa: E402
import cashflowparse  # noqa: E402
import budgetparse  # noqa: E402

from cryptography.hazmat.primitives.ciphers.aead import AESGCM

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
SITE = os.path.join(ROOT, "site")
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def now_iso():
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


# ---------- crypto ----------
def key():
    k = os.environ.get("DATA_KEY", "").strip()
    if not k:
        raise SystemExit("DATA_KEY secret is missing")
    raw = base64.b64decode(k)
    assert len(raw) == 32, "DATA_KEY must be 32 bytes, base64"
    return raw


def enc(obj, k):
    iv = os.urandom(12)
    ct = AESGCM(k).encrypt(iv, json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode(), None)
    return {"v": 1, "iv": base64.b64encode(iv).decode(), "ct": base64.b64encode(ct).decode()}


def dec(blob, k):
    return json.loads(AESGCM(k).decrypt(base64.b64decode(blob["iv"]), base64.b64decode(blob["ct"]), None))


def load_enc(path, k, default):
    try:
        with open(path) as f:
            return dec(json.load(f), k)
    except FileNotFoundError:
        return default


def save_enc(path, obj, k):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f:
        json.dump(enc(obj, k), f)


# ---------- download + xlsx -> grid ----------
def download(url=None, enc_name="sheet.enc"):
    test_file = os.environ.get({"cashflow_sheet.enc": "CF_FILE", "budget_sheet.enc": "BUDGET_FILE"}.get(enc_name, "SHEET_FILE"))
    if test_file:  # local testing
        return open(test_file, "rb").read(), None
    import requests
    url = url or (os.environ.get("SHEET_URL", "").strip() if enc_name == "sheet.enc" else "")
    if not url and os.path.exists(os.path.join(DATA, enc_name)):
        url = load_enc(os.path.join(DATA, enc_name), key(), {}).get("url", "")
    if not url:
        raise RuntimeError("SHEET_URL secret is not set yet, so the Excel file can't be downloaded")
    if "download=1" not in url:
        url += ("&" if "?" in url else "?") + "download=1"
    r = requests.get(url, timeout=60, headers={"User-Agent": "almentor-pmo-dashboard/1.0"})
    body = r.content
    if r.status_code != 200 or not body.startswith(b"PK"):
        raise RuntimeError(
            f"SharePoint returned HTTP {r.status_code} ({r.headers.get('content-type', '?')}) instead of the Excel file. "
            "The share link must be an 'Anyone with the link can view' link.")
    return body, r.headers.get("Last-Modified")


CUR_RE = re.compile(r'\[\$([^\]\-]+)|"([A-Za-z€$]{1,4})\s*"|([€$£])')


def cell_text(c):
    v = c.value
    if v is None:
        return ""
    fmt = c.number_format or ""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{MON[v.month - 1]}-{v.year % 100:02d}"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        if "%" in fmt:
            return f"{v * 100:.4f}%"
        s = f"{v:.6f}".rstrip("0").rstrip(".")
        m = CUR_RE.search(fmt)
        if m:
            cur = next(g for g in m.groups() if g).strip()
            cur = {"€": "€", "$": "$"}.get(cur, cur)
            return f"{cur} {s}" if cur not in ("$",) else f"${s}"
        return s
    return str(v)


def grid_from_xlsx(body, sheet=SHEET):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True, read_only=False)
    if sheet not in wb.sheetnames:
        raise RuntimeError(f'The workbook has no "{sheet}" tab (tabs: {", ".join(wb.sheetnames)})')
    ws = wb[sheet]
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        cells = [cell_text(c) for c in row]
        if any(x.strip() for x in cells):
            rows.append(cells)
    modified = wb.properties.modified
    return rows, (modified.replace(tzinfo=dt.timezone.utc).isoformat() if modified else None)


# ---------- PM changes ----------
STATUSES = ("active", "completed", "paused", "terminated")
PSEUDO_PM = {"closed projects": "completed", "terminated projects": "terminated"}


def migrate_state(state):
    """v1 {pms, assign:{project: pm}} -> v2 {team:[{name,title}], projects:{project:{members:[{name,pct}], status}}}.
    The placeholder PMs "Closed Projects" / "Terminated Projects" become project statuses."""
    if "team" in state:
        return state
    team = [{"name": n, "title": "Project Manager"} for n in state.get("pms", []) if n.lower() not in PSEUDO_PM]
    projects = {}
    for p, pm in state.get("assign", {}).items():
        st = PSEUDO_PM.get(pm.lower())
        projects[p] = {"members": [], "status": st} if st else {"members": [{"name": pm, "pct": 100}], "status": "active"}
        if not st and pm.lower() not in (t["name"].lower() for t in team):
            team.append({"name": pm, "title": "Project Manager"})
    return {"v": 2, "team": sorted(team, key=lambda t: t["name"].lower()), "projects": projects, "updatedAt": state.get("updatedAt")}


def apply_changes(state, raw):
    """Changes sent by the dashboard:
       {"team": {"set": [{"name","title"}], "remove": ["name"]},
        "projects": {"<project>": {"members": [{"name","pct"}] (optional), "status": "active|completed|paused|" (optional)}}}"""
    if not raw or not raw.strip():
        return False
    ch = json.loads(raw)
    before = json.dumps(state, sort_keys=True)
    clean = lambda s, n=60: re.sub(r"\s+", " ", str(s or "")).strip()[:n]
    team = [dict(t) for t in state.get("team", [])]
    idx = lambda n: next((i for i, t in enumerate(team) if t["name"].lower() == n.lower()), None)
    tc = ch.get("team") or {}
    projects = {k: dict(v) for k, v in state.get("projects", {}).items()}
    for r in tc.get("rename", [])[:20]:  # rename a person everywhere (team list and every project)
        old, to = clean(r.get("from")), clean(r.get("to"))
        if not old or not to or old.lower() == to.lower() and old == to:
            continue
        k, kt = idx(old), idx(to)
        if k is not None:
            if kt is not None and kt != k:
                team.pop(k)
            else:
                team[k]["name"] = to
        for p in projects.values():
            mem, merged = [], {}
            for m in p.get("members", []):
                n = to if m["name"].lower() == old.lower() else m["name"]
                if n.lower() in merged:
                    merged[n.lower()]["pct"] = min(100, merged[n.lower()]["pct"] + m["pct"])
                else:
                    merged[n.lower()] = {"name": n, "pct": m["pct"]}
                    mem.append(merged[n.lower()])
            p["members"] = mem
    for t in tc.get("set", [])[:100]:
        n, title = clean(t.get("name")), clean(t.get("title"), 80)
        if not n:
            continue
        k = idx(n)
        if k is None:
            team.append({"name": n, "title": title})
        else:
            team[k]["title"] = title or team[k].get("title", "")
    for n in tc.get("remove", [])[:100]:
        k = idx(clean(n))
        if k is not None:
            team.pop(k)
    for proj, pc in list((ch.get("projects") or {}).items())[:200]:
        proj = clean(proj, 200)
        if not proj or not isinstance(pc, dict):
            continue
        cur = dict(projects.get(proj, {"members": [], "status": ""}))
        if "members" in pc and isinstance(pc["members"], list):
            mem = []
            for m in pc["members"][:6]:
                n = clean(m.get("name"))
                try:
                    pct = max(0, min(100, int(round(float(m.get("pct", 0))))))
                except (TypeError, ValueError):
                    pct = 0
                if n and n.lower() not in (x["name"].lower() for x in mem):
                    mem.append({"name": n, "pct": pct})
                    if idx(n) is None:
                        team.append({"name": n, "title": ""})
            cur["members"] = mem
        if "status" in pc:
            st = clean(pc["status"], 20).lower()
            cur["status"] = st if st in STATUSES else ""
        if cur.get("members") or cur.get("status"):
            projects[proj] = cur
        else:
            projects.pop(proj, None)
    state["team"] = sorted(team, key=lambda t: t["name"].lower())
    state["projects"] = projects
    state["v"] = 2
    changed = json.dumps(state, sort_keys=True) != before
    if changed:
        state["updatedAt"] = now_iso()
    return changed


def strip_volatile(snap):
    return {k: v for k, v in (snap or {}).items() if k not in ("parsedAt",)}


def main():
    k = key()
    os.makedirs(SITE, exist_ok=True)
    status = {"checkedAt": now_iso(), "ok": True, "message": "", "requestId": os.environ.get("REQUEST_ID", "")}
    snap_path, state_path = os.path.join(DATA, "snapshot.enc"), os.path.join(DATA, "state.enc")
    old = load_enc(snap_path, k, None)
    state = load_enc(state_path, k, {"pms": [], "assign": {}})

    snap = old
    try:
        if os.environ.get("BOOTSTRAP"):
            status["checkedAt"] = old["parsedAt"]
            status["message"] = "Initial data"
            raise StopIteration
        body, last_mod_hdr = download()
        rows, modified = grid_from_xlsx(body)
        new = parse_grid(rows)
        bad = check(new)
        if not new["projects"] or not new["summary"]:
            raise RuntimeError("The sheet was read but no projects were found. Has the tab layout changed?")
        if len(bad) > max(3, len(new["sheetTotals"]) // 2):
            raise RuntimeError("Parsed numbers don't match the sheet's Totals row for: "
                               + ", ".join(b[0] for b in bad[:6]) + ". The tab layout may have changed.")
        new["checks"] = bad  # a few months off = the sheet's own Totals formula; shown in Findings
        new["sourceModified"] = modified or last_mod_hdr
        if old is None or strip_volatile(old) != strip_volatile(new):
            snap = new
            save_enc(snap_path, snap, k)
            status["message"] = "Updated from Excel"
        else:
            snap = old
            status["message"] = "No changes in Excel since last refresh"
        status["sourceModified"] = new["sourceModified"]
    except StopIteration:
        pass
    except Exception as e:  # keep last good data
        traceback.print_exc()
        status["ok"] = False
        status["message"] = str(e)[:400]
        if old is None:
            raise SystemExit("No data available yet: " + str(e))

    # ---- cash flow workbook ----
    cf_path = os.path.join(DATA, "cashflow.enc")
    cf_old = load_enc(cf_path, k, None)
    cf = cf_old
    cf_status = {"ok": True, "message": ""}
    if os.environ.get("BOOTSTRAP") or not os.path.exists(os.path.join(DATA, "cashflow_sheet.enc")):
        pass
    else:
        try:
            body, hdr = download(enc_name="cashflow_sheet.enc")
            rows, modified = grid_from_xlsx(body, cashflowparse.SHEET)
            new = cashflowparse.parse_grid(rows)
            if not new["cashIn"] or not new["cashOut"]:
                raise RuntimeError("Cash flow tab was read but no Cash In / Cash Out rows were found. Has the layout changed?")
            new["checks"] = cashflowparse.check(new)
            months_n = max(1, len(new["months"]))
            if len(new["checks"]) > months_n:  # layout problem, not a few stray cells
                raise RuntimeError("Cash flow rows don't add up to the sheet's totals in most months. The tab layout may have changed.")
            new["sourceModified"] = modified or hdr
            if cf_old is None or strip_volatile(cf_old) != strip_volatile(new):
                cf = new
                save_enc(cf_path, cf, k)
        except Exception as e:
            traceback.print_exc()
            cf_status = {"ok": False, "message": "Cash flow: " + str(e)[:300]}
    if cf is not None and cf_status["ok"] and cf is not cf_old:
        cf["parsedAt"] = status["checkedAt"]
    status["cashflow"] = cf_status
    if not cf_status["ok"]:
        status["ok"] = False
        status["message"] = (status["message"] + " · " if status["message"] else "") + cf_status["message"]

    # ---- budget & cost control workbook ----
    bd_path = os.path.join(DATA, "budget.enc")
    bd_old = load_enc(bd_path, k, None)
    bd = bd_old
    bd_status = {"ok": True, "message": ""}
    if not os.environ.get("BOOTSTRAP") and os.path.exists(os.path.join(DATA, "budget_sheet.enc")):
        try:
            import openpyxl
            body, hdr = download(enc_name="budget_sheet.enc")
            wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
            grids = {n: [[cell_text(c) for c in row] for row in wb[n].iter_rows(min_row=1, max_row=wb[n].max_row, min_col=1, max_col=wb[n].max_column)] for n in wb.sheetnames}
            new = budgetparse.parse_workbook(grids, {n: wb[n].sheet_state for n in wb.sheetnames})
            if not new["projects"]:
                raise RuntimeError("Budget workbook was read but no project sheets could be parsed. " + "; ".join(new["problems"][:3]))
            m = wb.properties.modified
            new["sourceModified"] = m.replace(tzinfo=dt.timezone.utc).isoformat() if m else hdr
            if bd_old is None or strip_volatile(bd_old) != strip_volatile(new):
                bd = new
                save_enc(bd_path, bd, k)
        except Exception as e:
            traceback.print_exc()
            bd_status = {"ok": False, "message": "Budget: " + str(e)[:300]}
    if bd is not None and bd_status["ok"] and bd is not bd_old:
        bd["parsedAt"] = status["checkedAt"]
    status["budget"] = bd_status
    if not bd_status["ok"]:
        status["ok"] = False
        status["message"] = (status["message"] + " · " if status["message"] else "") + bd_status["message"]

    migrated = "team" not in state
    state = migrate_state(state)
    if migrated:
        save_enc(state_path, state, k)

    if apply_changes(state, os.environ.get("CHANGES", "")):
        save_enc(state_path, state, k)

    snap = dict(snap)
    if status["ok"] and not os.environ.get("BOOTSTRAP"):
        snap["parsedAt"] = status["checkedAt"]
    bundle = {"snapshot": snap, "cashflow": cf, "budget": bd, "state": state, "status": status,
              "repo": os.environ.get("GITHUB_REPOSITORY", ""),
              "token": os.environ.get("DISPATCH_TOKEN", "").strip()
                       or load_enc(os.path.join(DATA, "token.enc"), k, {}).get("token", "")}
    with open(os.path.join(SITE, "data.enc"), "w") as f:
        json.dump(enc(bundle, k), f)
    with open(os.path.join(SITE, "status.json"), "w") as f:
        json.dump({**status, "dataHash": hashlib.sha256(json.dumps(strip_volatile(snap), sort_keys=True).encode()).hexdigest()[:16]}, f)
    for fn in ("index.html", "keys.json", "favicon.png"):
        if os.path.exists(os.path.join(ROOT, fn)):
            shutil.copy(os.path.join(ROOT, fn), os.path.join(SITE, fn))
    if os.path.isdir(os.path.join(ROOT, "vendor")):
        shutil.copytree(os.path.join(ROOT, "vendor"), os.path.join(SITE, "vendor"), dirs_exist_ok=True)
    open(os.path.join(SITE, ".nojekyll"), "w").close()
    print(json.dumps(status))
    if not status["ok"]:
        # data stayed published, but flag the run so the Actions tab shows it
        with open(os.environ.get("GITHUB_STEP_SUMMARY", os.devnull), "a") as f:
            f.write(f"**Sync failed:** {status['message']}\n")


if __name__ == "__main__":
    main()
