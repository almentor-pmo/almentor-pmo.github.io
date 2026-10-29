#!/usr/bin/env python3
"""Hourly sync for the almentor Quality (QA) dashboard  ->  https://almentor-pmo.github.io/quality/

1. Downloads the Quality Tracking workbook from its OneDrive/SharePoint "anyone" link
   (stored encrypted in data/qa/sheet.enc) and reads the Courses, Issue Log and Lists tabs.
2. Writes quality/data.enc (AES-256-GCM, QA key) and quality/status.json for the page.

Keys: the repo secret DATA_KEY decrypts data/qa/key.enc, which holds the QA dashboard's own
key. QA logins (quality/keys.json) unwrap only the QA key, so they can't open the PMO data.
If the download or parsing fails, the last good data is kept and the error is shown on the page.
Local test:  QA_FILE=book.xlsx DATA_KEY=... python scripts/qa_sync.py
"""
import base64, datetime as dt, hashlib, io, json, os, re, sys, traceback

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sync import enc, load_enc, save_enc, now_iso, key as data_key  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
QDATA = os.path.join(ROOT, "data", "qa")
SITE = os.path.join(ROOT, "quality")


def qa_key():
    blob = load_enc(os.path.join(QDATA, "key.enc"), data_key(), None)
    if not blob:
        raise SystemExit("data/qa/key.enc is missing")
    return base64.b64decode(blob["key"])


def download():
    f = os.environ.get("QA_FILE")
    if f:
        return open(f, "rb").read(), None
    import requests
    url = load_enc(os.path.join(QDATA, "sheet.enc"), data_key(), {}).get("url", "")
    if not url:
        raise RuntimeError("The Quality Tracking sheet link isn't set yet")
    if "download=1" not in url:
        url += ("&" if "?" in url else "?") + "download=1"
    r = requests.get(url, timeout=60, headers={"User-Agent": "almentor-qa-dashboard/1.0"})
    if r.status_code != 200 or not r.content.startswith(b"PK"):
        raise RuntimeError(f"OneDrive returned HTTP {r.status_code} ({r.headers.get('content-type', '?')}) instead of the Excel file. "
                           "The share link must be an 'Anyone with the link can view' link.")
    return r.content, r.headers.get("Last-Modified")


# ---------- reading ----------
def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # Excel serial
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = txt(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%y", "%d-%b-%Y", "%d %b %Y", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def header_row(ws, must, scan=12):
    for r in range(1, scan + 1):
        vals = [txt(c.value).lower() for c in ws[r]]
        if all(any(v.startswith(m) for v in vals) for m in must):
            return r, vals
    raise RuntimeError(f'Couldn\'t find the header row on the "{ws.title}" tab (looking for {", ".join(must)}). Has the layout changed?')


def col(vals, *names):
    for n in names:
        for i, v in enumerate(vals):
            if v == n or v.startswith(n):
                return i
    return None


def parse(body):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    logtab = "Issue Log" if "Issue Log" in wb.sheetnames else "Review Log"
    for t in ("Courses", logtab, "Lists"):
        if t not in wb.sheetnames:
            raise RuntimeError(f'The workbook has no "{t}" tab (tabs: {", ".join(wb.sheetnames)})')
    problems = []

    # Lists
    ls = wb["Lists"]
    hr, hv = header_row(ls, ["stage", "responsible team"])
    grid = [[c.value for c in row] for row in ls.iter_rows(min_row=hr + 1, max_row=ls.max_row)]
    def colvals(i):
        return [txt(r[i]) for r in grid if i is not None and i < len(r) and txt(r[i])]
    ci = col(hv, "stage"); ti = col(hv, "responsible team")
    stages = []
    for r in grid:
        s, t = txt(r[ci]) if ci < len(r) else "", txt(r[ti]) if ti < len(r) else ""
        if s:
            stages.append({"name": s, "team": t})
    teams = colvals(col(hv, "teams"))
    sev_i = col(hv, "severity")
    severity = [{"name": txt(r[sev_i]), "desc": txt(r[sev_i + 1]) if sev_i + 1 < len(r) else ""} for r in grid if sev_i is not None and txt(r[sev_i])]
    reviewers = colvals(col(hv, "qa reviewers"))
    projects = colvals(col(hv, "projects"))
    stage_names = [s["name"] for s in stages]
    cats = {}
    for i, v in enumerate([txt(c.value) for c in ls[hr]]):
        if v in stage_names and i != ci:
            cats[v] = colvals(i)

    # Courses
    cs = wb["Courses"]
    hr, hv = header_row(cs, ["project", "course"])
    hraw = [txt(c.value) for c in cs[hr]]
    ip, it, ic, ir, isd = col(hv, "project"), col(hv, "type"), col(hv, "course"), col(hv, "qa reviewer"), col(hv, "start date")
    stage_cols = [(i, h) for i, h in enumerate(hraw) if h in stage_names]
    courses = []
    for n, row in enumerate(cs.iter_rows(min_row=hr + 1, max_row=cs.max_row, values_only=True), start=hr + 1):
        name = txt(row[ic]) if ic is not None and ic < len(row) else ""
        proj = txt(row[ip]) if ip is not None and ip < len(row) else ""
        if not name and not proj:
            continue
        courses.append({"row": n, "project": proj, "type": txt(row[it]) if it is not None else "", "course": name,
                        "reviewer": txt(row[ir]) if ir is not None else "",
                        "start": as_date(row[isd]) if isd is not None else None,
                        "stages": {h: txt(row[i]) for i, h in stage_cols if i < len(row)}})

    # Review Log
    lg = wb[logtab]
    hr, hv = header_row(lg, ["date logged", "stage", "severity"])
    m = {k: col(hv, *names) for k, names in {
        "date": ("date logged",), "course": ("course",), "stage": ("stage",), "reviewer": ("qa reviewer",),
        "category": ("issue category", "error category"), "desc": ("description",), "severity": ("severity",), "round": ("review round",),
        "status": ("status",), "fixed": ("date fixed",), "override": ("repeated (override)",), "notes": ("notes",),
        "auto": ("repeated (auto)",), "final": ("repeated (final)",)}.items()}
    log = []
    for n, row in enumerate(lg.iter_rows(min_row=hr + 1, max_row=lg.max_row, values_only=True), start=hr + 1):
        g = lambda k: row[m[k]] if m[k] is not None and m[k] < len(row) else None
        if not txt(g("course")) and not txt(g("stage")) and not txt(g("category")) and not txt(g("desc")):
            continue
        rd = g("round")
        try:
            rd = int(float(rd)) if txt(rd) else None
        except ValueError:
            rd = None
        log.append({"row": n, "date": as_date(g("date")), "course": txt(g("course")), "stage": txt(g("stage")),
                    "reviewer": txt(g("reviewer")), "category": txt(g("category")), "desc": txt(g("desc")),
                    "severity": txt(g("severity")), "round": rd, "status": txt(g("status")), "fixed": as_date(g("fixed")),
                    "override": txt(g("override")), "notes": txt(g("notes")),
                    "sheetFinal": txt(g("final"))})
    m_ = wb.properties.modified
    return {"stages": stages, "teams": teams, "severity": severity, "reviewers": reviewers, "projects": projects,
            "categories": cats, "courses": courses, "log": log, "problems": problems,
            "sourceModified": m_.replace(tzinfo=dt.timezone.utc).isoformat() if m_ else None}


def main():
    k = qa_key()
    os.makedirs(SITE, exist_ok=True)
    snap_path = os.path.join(QDATA, "snapshot.enc")
    old = load_enc(snap_path, k, None)
    status = {"checkedAt": now_iso(), "ok": True, "message": "", "requestId": os.environ.get("REQUEST_ID", "")}
    snap = old
    try:
        body, hdr = download()
        new = parse(body)
        if not new["stages"]:
            raise RuntimeError("The Lists tab has no stages. Has the layout changed?")
        new["sourceModified"] = new["sourceModified"] or hdr
        strip = lambda s: {x: y for x, y in (s or {}).items() if x != "parsedAt"}
        if old is None or strip(old) != strip(new):
            snap = new
            snap["parsedAt"] = status["checkedAt"]
            save_enc(snap_path, snap, k)
            status["message"] = "Updated from Excel"
        else:
            status["message"] = "No changes in Excel since last refresh"
        status["sourceModified"] = new["sourceModified"]
    except Exception as e:
        traceback.print_exc()
        status["ok"] = False
        status["message"] = str(e)[:400]
        if old is None:
            raise SystemExit("No QA data available yet: " + str(e))
    dk = data_key()
    bundle = {"data": snap, "status": status, "repo": os.environ.get("GITHUB_REPOSITORY", "almentor-pmo/almentor-pmo.github.io"),
              "token": load_enc(os.path.join(ROOT, "data", "token.enc"), dk, {}).get("token", "")}
    with open(os.path.join(SITE, "data.enc"), "w") as f:
        json.dump(enc(bundle, k), f)
    with open(os.path.join(SITE, "status.json"), "w") as f:
        json.dump({**status, "dataHash": hashlib.sha256(json.dumps(snap, sort_keys=True, default=str).encode()).hexdigest()[:16]}, f)
    print(json.dumps(status))
    if not status["ok"]:
        with open(os.environ.get("GITHUB_STEP_SUMMARY", os.devnull), "a") as f:
            f.write(f"**QA sync failed:** {status['message']}\n")


if __name__ == "__main__":
    main()
